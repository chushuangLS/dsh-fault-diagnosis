"""
生成课表解析器的测试样本（.xls / .xlsx）

为什么要生成而不是直接用真实的课表文件：
  真实文件里有教师姓名、教学班、选课人数，
  提交进仓库就泄露了（比赛要求不得出现真实教师姓名）。
  所以这里用完全虚构的课程和教师名造一份结构相同的样本。

表格结构照抄教务系统的导出格式（实测样本：第一实验楼机房C104.xls）：

  第 0 行   学期 | 教室场地课表 | 座位数
  第 1 行   节次 | (空) | 星期一 … 星期日
  第 2-13 行  上下午晚上 × 星期
  第 14 行  注释（含学期起止与总周数）

  格子格式：课程名/(N-N节)X-Y周/教师/教学班/选课人数：N

重新生成：
    python tools/make_schedule_fixture.py
"""

import os
import sys

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(HERE, 'tests', 'fixtures')

# 虚构数据：课程、教师、教学班全部是编的
COURSES = {
    # (星期, 节次) -> 该格的课程列表
    (1, 1): [
        ('测试课程甲/(1-1节)1-12周/张甲/测试课程甲-0001/选课人数：50', 1),
    ],
    (2, 1): [
        ('测试课程甲/(1-1节)1-12周/张甲/测试课程甲-0001/选课人数：50', 1),
        ('测试课程乙/(1-1节)1-12周/李乙/测试课程乙-0002/选课人数：60', 1),
    ],
    (3, 2): [
        ('测试课程丙/(2-2节)1-16周/王丙/测试课程丙-0003/选课人数：45', 1),
    ],
    # 同一格两门课，不同周次交替 —— 真实数据里就有这种
    (4, 3): [
        ('测试课程丁/戊/(3-3节)1-8周/赵丁/测试课程丁/戊-0004/选课人数：30', 1),
        ('测试课程己/(3-3节)9-16周/钱己/测试课程己-0005/选课人数：40', 1),
    ],
    (5, 9): [
        ('测试课程庚/(9-9节)1-6周/孙庚/测试课程庚-0006/选课人数：35', 1),
    ],
}

PERIOD_NAMES = ['一', '二', '三', '四', '五', '六',
                '七', '八', '九', '十', '十一', '十二']
SECTION_OF = {1: '上午', 2: '上午', 3: '上午', 4: '上午',
              5: '下午', 6: '下午', 7: '下午', 8: '下午',
              9: '晚上', 10: '晚上', 11: '晚上', 12: '晚上'}
WEEKDAYS = ['星期一', '星期二', '星期三', '星期四', '星期五', '星期六', '星期日']

ROOM_LABEL = '第一测试楼机房T101'
SEATS = 60
NOTE = ('注--内容顺序为：课程<>周次<>教师<>教学班<>选课人数          '
        '本学期2026-09-07正式上课至2027-01-24结束，共20周.')


def build_grid():
    """拼出表格的二维内容"""
    grid = []

    # 第 0 行：标题
    grid.append(['2026-2027学年第1学期', '2026-2027学年第1学期',
                 ROOM_LABEL + '场地课表', ROOM_LABEL + '场地课表',
                 ROOM_LABEL + '场地课表',
                 f'座位数：{SEATS}', f'座位数：{SEATS}'])

    # 第 1 行：表头
    grid.append(['节次', ''] + WEEKDAYS)

    # 第 2-13 行：节次
    for period in range(1, 13):
        row = ['', PERIOD_NAMES[period - 1]]
        for day in range(1, 8):
            cells = COURSES.get((day, period))
            row.append('\n'.join(c[0] for c in cells) if cells else '')
        grid.append(row)

    # 最后一行：注释
    grid.append([NOTE])

    return grid


def write_xls(path, grid):
    import xlwt

    book = xlwt.Workbook(encoding='utf-8')
    sh = book.add_sheet('Sheet0')
    for r, row in enumerate(grid):
        for c, val in enumerate(row):
            sh.write(r, c, val)
    book.save(path)


def write_xlsx(path, grid):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet0'
    for r, row in enumerate(grid, start=1):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)
    wb.save(path)


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    grid = build_grid()

    xls = os.path.join(OUT_DIR, 'schedule_sample.xls')
    xlsx = os.path.join(OUT_DIR, 'schedule_sample.xlsx')

    write_xls(xls, grid)
    write_xlsx(xlsx, grid)

    print(f'已生成 {xls}')
    print(f'已生成 {xlsx}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
